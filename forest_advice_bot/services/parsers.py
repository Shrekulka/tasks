# forest-advice-bot/services/parsers.py

"""Разбор PDF, DOCX и TXT в сегменты текста с номером страницы или заголовком раздела.

Формат определяется по расширению И по содержимому: расширению самому по себе не верим.
Новый формат подключается одной записью FormatSpec в _FORMATS, остальной код не меняется.
"""

import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath

from docx import Document as open_docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PasswordType, PdfReader

from config_data.config import RagConfig
from errors.errors import EmptyDocument, UnreadableDocument, UnsupportedFormat

_MB = 1024 * 1024
_PDF_SIGNATURE = b"%PDF-"
_ZIP_SIGNATURE = b"PK\x03\x04"
_DOCX_MAIN_PART = "word/document.xml"
_TEXT_SNIFF_BYTES = 8192
_TEXT_ENCODINGS = ("utf-8-sig", "cp1251")  # сначала UTF-8, затем старые украинские тексты
_HEADING_RE = re.compile(r"^(heading|заголовок)\s*\d*$|^title$", re.IGNORECASE)
_SPACES_RE = re.compile(r"[ \t\u00a0]+")
_BLANK_LINES_RE = re.compile(r"\n{3,}")


@dataclass(frozen=True, slots=True)
class Segment:
    """Кусок текста документа и его «адрес» для цитирования."""

    text: str
    page: int | None = None
    section: str | None = None


@dataclass(frozen=True, slots=True)
class ParseLimits:
    """Защита от тяжёлых и вредоносных файлов. Значения приходят из конфигурации."""

    max_pdf_pages: int
    max_zip_entries: int
    max_unpacked_bytes: int

    @classmethod
    def from_config(cls, rag: RagConfig) -> "ParseLimits":
        return cls(
            max_pdf_pages=rag.max_pdf_pages,
            max_zip_entries=rag.max_zip_entries,
            max_unpacked_bytes=rag.max_unpacked_mb * _MB,
        )


@dataclass(frozen=True, slots=True)
class FormatSpec:
    name: str
    extensions: frozenset[str]
    matches: Callable[[bytes], bool]
    parse: Callable[[bytes, ParseLimits], list[Segment]]


# ---------------------------------------------------------------------- общее


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _SPACES_RE.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANK_LINES_RE.sub("\n\n", text).strip()


def _require_text(segments: list[Segment]) -> list[Segment]:
    if not segments:
        raise EmptyDocument("в файле нет текста")
    return segments


# ---------------------------------------------------------------------- PDF


def _is_pdf(data: bytes) -> bool:
    return data.startswith(_PDF_SIGNATURE)


def _parse_pdf(data: bytes, limits: ParseLimits) -> list[Segment]:
    segments: list[Segment] = []
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted and reader.decrypt("") == PasswordType.NOT_DECRYPTED:
            raise UnreadableDocument("PDF защищён паролем")
        if len(reader.pages) > limits.max_pdf_pages:
            raise UnreadableDocument(f"в PDF больше {limits.max_pdf_pages} страниц")
        for number, page in enumerate(reader.pages, start=1):
            text = _clean(page.extract_text() or "")
            if text:
                segments.append(Segment(text=text, page=number))
    except (UnreadableDocument, EmptyDocument):
        raise
    except Exception as exc:  # noqa: BLE001  pypdf бросает десятки разных исключений на битых файлах
        raise UnreadableDocument(f"PDF не читается: {type(exc).__name__}") from exc
    return _require_text(segments)


# ---------------------------------------------------------------------- DOCX


def _is_docx(data: bytes) -> bool:
    if not data.startswith(_ZIP_SIGNATURE):
        return False
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            return _DOCX_MAIN_PART in archive.namelist()
    except zipfile.BadZipFile:
        return False


def _check_zip_limits(data: bytes, limits: ParseLimits) -> None:
    """Архивная бомба: смотрим на заявленные размеры до того, как что-то распаковывать."""
    with zipfile.ZipFile(BytesIO(data)) as archive:
        infos = archive.infolist()
    if len(infos) > limits.max_zip_entries:
        raise UnreadableDocument("слишком много записей в архиве DOCX")
    if sum(info.file_size for info in infos) > limits.max_unpacked_bytes:
        raise UnreadableDocument("DOCX слишком большой после распаковки")


def _table_rows(table: Table) -> list[str]:
    rows: list[str] = []
    for row in table.rows:
        seen: set[int] = set()
        cells: list[str] = []
        for cell in row.cells:
            if id(cell._tc) in seen:  # объединённая ячейка повторяется в row.cells
                continue
            seen.add(id(cell._tc))
            text = _clean(cell.text).replace("\n", " ")
            if text:
                cells.append(text)
        if cells:
            rows.append(" | ".join(cells))
    return rows


def _parse_docx(data: bytes, limits: ParseLimits) -> list[Segment]:
    _check_zip_limits(data, limits)
    try:
        document = open_docx(BytesIO(data))
        segments: list[Segment] = []
        section: str | None = None
        # Обходим тело документа по порядку: абзацы и таблицы идут как в исходнике.
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                paragraph = Paragraph(child, document)
                text = _clean(paragraph.text)
                if not text:
                    continue
                style = paragraph.style.name if paragraph.style is not None else ""
                if _HEADING_RE.match(style or ""):
                    section = text
                segments.append(Segment(text=text, section=section))
            elif child.tag == qn("w:tbl"):
                segments.extend(
                    Segment(text=row, section=section)
                    for row in _table_rows(Table(child, document))
                )
    except UnreadableDocument:
        raise
    except Exception as exc:  # noqa: BLE001  python-docx и lxml бросают разные исключения
        raise UnreadableDocument(f"DOCX не читается: {type(exc).__name__}") from exc
    return _require_text(segments)


# ---------------------------------------------------------------------- TXT


def _is_text(data: bytes) -> bool:
    head = data[:_TEXT_SNIFF_BYTES]
    return not (
        head.startswith(_PDF_SIGNATURE) or head.startswith(_ZIP_SIGNATURE) or b"\x00" in head
    )


def _parse_txt(data: bytes, limits: ParseLimits) -> list[Segment]:
    del limits  # размер текста уже ограничен лимитом файла
    for encoding in _TEXT_ENCODINGS:
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise UnreadableDocument("не удалось определить кодировку текста")
    cleaned = _clean(text)
    return _require_text([Segment(text=cleaned)] if cleaned else [])


# ---------------------------------------------------------------------- реестр

_FORMATS: tuple[FormatSpec, ...] = (
    FormatSpec("pdf", frozenset({".pdf"}), _is_pdf, _parse_pdf),
    FormatSpec("docx", frozenset({".docx"}), _is_docx, _parse_docx),
    FormatSpec("txt", frozenset({".txt"}), _is_text, _parse_txt),
)


def supported_extensions() -> tuple[str, ...]:
    """Для текстов бота («надішли PDF, Word або TXT») и проверки до скачивания."""
    return tuple(sorted(ext for spec in _FORMATS for ext in spec.extensions))


def detect_format(filename: str, data: bytes) -> FormatSpec:
    """Формат по расширению, подтверждённый содержимым.

    UnsupportedFormat(reason="extension"): расширение не из списка.
    UnsupportedFormat(reason="mismatch"): содержимое не соответствует расширению.
    """
    extension = PurePath(filename).suffix.lower()
    spec = next((s for s in _FORMATS if extension in s.extensions), None)
    if spec is None:
        raise UnsupportedFormat("extension", extension or "без расширения")
    if not spec.matches(data):
        raise UnsupportedFormat("mismatch", f"{extension} не соответствует содержимому")
    return spec


def extract_segments(
    filename: str, data: bytes, limits: ParseLimits
) -> tuple[FormatSpec, list[Segment]]:
    """Определяет формат и возвращает сегменты текста.

    Бросает UnsupportedFormat, EmptyDocument или UnreadableDocument.
    """
    spec = detect_format(filename, data)
    return spec, spec.parse(data, limits)
